"""Small in-memory fixtures: real PDF/DOCX/XLSX bytes without files on disk."""

import io

import docx
import openpyxl


def make_pdf(pages):
    """Minimal valid PDF, one Helvetica text line per list item, one page per entry."""
    objects = []

    def add(body):
        objects.append(body)
        return len(objects)

    font = add(b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>')
    page_ids = []
    pages_id = len(pages) * 2 + 2  # reserved slot, filled below
    for lines in pages:
        stream = ['BT', '/F1 11 Tf', '14 TL', '50 780 Td']
        for line in lines:
            escaped = line.replace('\\', '\\\\').replace('(', '\\(').replace(')', '\\)')
            stream.append(f'({escaped}) Tj T*')
        stream.append('ET')
        data = '\n'.join(stream).encode('latin-1')
        content = add(b'<< /Length %d >>\nstream\n' % len(data) + data + b'\nendstream')
        page_ids.append(
            add(
                b'<< /Type /Page /Parent %d 0 R /MediaBox [0 0 612 842] '
                b'/Resources << /Font << /F1 %d 0 R >> >> /Contents %d 0 R >>'
                % (pages_id, font, content)
            )
        )
    kids = b' '.join(b'%d 0 R' % pid for pid in page_ids)
    assert add(b'<< /Type /Pages /Kids [%s] /Count %d >>' % (kids, len(page_ids))) == pages_id
    catalog = add(b'<< /Type /Catalog /Pages %d 0 R >>' % pages_id)

    out = io.BytesIO()
    out.write(b'%PDF-1.4\n')
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b'%d 0 obj\n' % number + body + b'\nendobj\n')
    xref = out.tell()
    out.write(b'xref\n0 %d\n0000000000 65535 f \n' % (len(objects) + 1))
    for offset in offsets:
        out.write(b'%010d 00000 n \n' % offset)
    out.write(
        b'trailer\n<< /Size %d /Root %d 0 R >>\nstartxref\n%d\n%%%%EOF\n'
        % (len(objects) + 1, catalog, xref)
    )
    return out.getvalue()


def make_docx():
    document = docx.Document()
    document.core_properties.title = 'Hostel Rules'
    document.add_heading('Hostel Rules', level=0)
    document.add_heading('Timings', level=1)
    document.add_paragraph('Gates close at 10:30 pm for first year students.')
    document.add_paragraph('Visitors must sign the register.', style='List Bullet')
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = 'Hostel', 'Fee'
    table.cell(1, 0).text, table.cell(1, 1).text = 'Hall A', '1,20,000'
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_xlsx():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'Fees 2026-27'
    sheet.append(['Programme', 'Tuition | per sem'])
    sheet.append(['BE COE', 225000])
    sheet.append([None, None])
    sheet.append(['BE ECE', 210000])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
