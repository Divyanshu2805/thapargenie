"""Turn source files into Markdown.

Every extractor returns Markdown in which page boundaries (when the source has pages)
are marked by lines of the form `[[page N]]`. The chunker reads and strips them.
"""

import csv
import io
import re
from collections import Counter
from dataclasses import dataclass

import docx
import openpyxl
import trafilatura
from common.text import normalize_text
from docx.table import Table
from docx.text.paragraph import Paragraph
from pypdf import PdfReader, PdfWriter
from pypdf.errors import PdfReadError

from knowledge.ingest.filetypes import UnsupportedFile, decode_text
from knowledge.models import SourceType

MIN_CHARS_PER_PAGE = 200
SMART_WINDOW_PAGES = 20
MAX_SHEET_ROWS = 5000
EDGE_LINES = 2
MAX_EDGE_LINE_CHARS = 120
REPEAT_RATIO = 0.6

SMART_PROMPT = """Convert pages {first}-{last} of this PDF (pages are numbered from {first}) \
into clean Markdown.

Rules:
- Before the content of each page, write a line containing only [[page N]] with its page number.
- Keep all text. Do not summarise, translate, correct or add anything.
- Use #, ##, ### for headings as they appear.
- Reproduce every table as a GitHub-flavoured Markdown table with its header row. \
Merged cells: repeat the value.
- Drop page headers, footers, page numbers and watermarks that repeat on every page.
- Output only the Markdown."""


@dataclass
class Extracted:
    markdown: str
    page_count: int | None = None
    title: str = ''
    used_smart: bool = False


def page_marker(number):
    return f'[[page {number}]]'


def markdown_table(rows):
    """Rows of cell values -> GitHub table. The first row is the header."""
    cleaned = [
        [re.sub(r'\s+', ' ', str(cell if cell is not None else '')).replace('|', '\\|').strip()
         for cell in row]
        for row in rows
    ]
    cleaned = [row for row in cleaned if any(row)]
    if not cleaned:
        return ''
    width = max(len(row) for row in cleaned)
    cleaned = [row + [''] * (width - len(row)) for row in cleaned]
    header, *body = cleaned
    lines = ['| ' + ' | '.join(header) + ' |', '|' + '---|' * width]
    lines += ['| ' + ' | '.join(row) + ' |' for row in body]
    return '\n'.join(lines)


# -- PDF -----------------------------------------------------------------------------


def open_pdf(data):
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted and not reader.decrypt(''):
            raise UnsupportedFile('Password-protected PDFs are not supported.')
        return reader
    except PdfReadError as exc:
        raise UnsupportedFile('The PDF could not be read.') from exc


def _strip_repeated_edges(pages):
    """Remove header/footer lines that repeat on most pages.

    Only short lines at the very top or bottom of a page qualify, and a page is never
    emptied: body text that happens to repeat must survive.
    """
    if len(pages) < 3:
        return pages
    counts = Counter()
    for lines in pages:
        edges = {
            line
            for line in lines[:EDGE_LINES] + lines[-EDGE_LINES:]
            if line and len(line) <= MAX_EDGE_LINE_CHARS
        }
        counts.update(edges)
    repeated = {line for line, count in counts.items() if count / len(pages) >= REPEAT_RATIO}
    cleaned = []
    for lines in pages:
        edge_positions = set(range(EDGE_LINES)) | set(range(len(lines) - EDGE_LINES, len(lines)))
        kept = [
            line
            for position, line in enumerate(lines)
            if not (position in edge_positions and line in repeated)
        ]
        cleaned.append(kept if any(kept) else lines)
    return cleaned


def pdf_text_pages(reader):
    pages = []
    for page in reader.pages:
        try:
            text = page.extract_text() or ''
        except Exception:  # a single broken page should not sink the document
            text = ''
        pages.append([line.strip() for line in normalize_text(text).split('\n')])
    return _strip_repeated_edges(pages)


def needs_smart_parsing(pages):
    """Scanned or garbled PDFs have little usable text in their text layer."""
    if not pages:
        return True
    chars = [sum(len(line) for line in lines) for lines in pages]
    thin_pages = sum(1 for count in chars if count < MIN_CHARS_PER_PAGE)
    return thin_pages / len(pages) > 0.5


def _fast_pdf(pages):
    parts = []
    for number, lines in enumerate(pages, start=1):
        body = '\n'.join(lines).strip()
        if body:
            parts.append(f'{page_marker(number)}\n{body}')
    return '\n\n'.join(parts)


def _pdf_window(reader, start, stop):
    writer = PdfWriter()
    for index in range(start, stop):
        writer.add_page(reader.pages[index])
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _smart_pdf(reader, llm):
    total = len(reader.pages)
    parts = []
    for start in range(0, total, SMART_WINDOW_PAGES):
        stop = min(start + SMART_WINDOW_PAGES, total)
        prompt = SMART_PROMPT.format(first=start + 1, last=stop)
        parts.append(llm.read_pdf(_pdf_window(reader, start, stop), prompt).strip())
    return '\n\n'.join(parts)


def extract_pdf(data, *, mode='auto', llm=None, max_pages=None):
    reader = open_pdf(data)
    page_count = len(reader.pages)
    if max_pages and page_count > max_pages:
        raise UnsupportedFile(f'The PDF has {page_count} pages; the limit is {max_pages}.')
    title = (reader.metadata.title if reader.metadata else None) or ''
    if mode == 'smart':
        use_smart = True
    else:
        pages = pdf_text_pages(reader)
        use_smart = mode == 'auto' and needs_smart_parsing(pages)
    if use_smart:
        if llm is None:
            raise UnsupportedFile('This PDF needs smart parsing, which is not configured.')
        markdown = _smart_pdf(reader, llm)
    else:
        markdown = _fast_pdf(pages)
    return Extracted(normalize_text(markdown), page_count, str(title).strip(), use_smart)


# -- Office / text ------------------------------------------------------------------


def _docx_blocks(document):
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit('}', 1)[-1]
        if tag == 'p':
            yield Paragraph(child, document)
        elif tag == 'tbl':
            yield Table(child, document)


def extract_docx(data):
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise UnsupportedFile('The Word document could not be read.') from exc
    parts = []
    for block in _docx_blocks(document):
        if isinstance(block, Table):
            rows = [[cell.text for cell in row.cells] for row in block.rows]
            table = markdown_table(rows)
            if table:
                parts.append(table)
            continue
        text = block.text.strip()
        if not text:
            continue
        style = (block.style.name if block.style is not None else '').lower()
        level = re.match(r'heading (\d)', style)
        if style == 'title':
            parts.append(f'# {text}')
        elif level:
            parts.append('#' * min(int(level.group(1)) + 1, 6) + f' {text}')
        elif 'list' in style:
            parts.append(f'- {text}')
        else:
            parts.append(text)
    title = document.core_properties.title or ''
    return Extracted(normalize_text('\n\n'.join(parts)), title=title.strip())


def extract_xlsx(data):
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise UnsupportedFile('The Excel file could not be read.') from exc
    parts = []
    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            if any(cell not in (None, '') for cell in row):
                rows.append(list(row))
            if len(rows) >= MAX_SHEET_ROWS:
                break
        table = markdown_table(rows)
        if table:
            parts.append(f'## {sheet.title}\n\n{table}')
    workbook.close()
    return Extracted(normalize_text('\n\n'.join(parts)))


def extract_csv(data):
    text = decode_text(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=',;\t|')
    except csv.Error:
        dialect = csv.excel
    rows = []
    for row in csv.reader(io.StringIO(text), dialect):
        rows.append(row)
        if len(rows) >= MAX_SHEET_ROWS:
            break
    return Extracted(markdown_table(rows))


def extract_html(html, url=None):
    markdown = trafilatura.extract(
        html,
        url=url,
        output_format='markdown',
        include_tables=True,
        include_links=False,
        include_images=False,
        favor_recall=True,
    )
    if not markdown:
        raise UnsupportedFile('No readable content found on the page.')
    metadata = trafilatura.extract_metadata(html, default_url=url)
    title = (metadata.title if metadata else '') or ''
    return Extracted(normalize_text(markdown), title=title.strip())


def extract(data, source_type, *, mode='auto', llm=None, max_pages=None, url=None):
    if source_type == SourceType.PDF:
        return extract_pdf(data, mode=mode, llm=llm, max_pages=max_pages)
    if source_type == SourceType.DOCX:
        return extract_docx(data)
    if source_type == SourceType.XLSX:
        return extract_xlsx(data)
    if source_type == SourceType.CSV:
        return extract_csv(data)
    if source_type in (SourceType.HTML, SourceType.URL):
        return extract_html(decode_text(data), url=url)
    if source_type == SourceType.TEXT:
        return Extracted(normalize_text(decode_text(data)))
    raise UnsupportedFile(f'Cannot extract {source_type}.')
